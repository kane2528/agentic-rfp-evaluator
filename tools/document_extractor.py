"""Text extraction for supported proposal documents."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from pptx import Presentation

from tools.pdf_extractor import PDFExtractionError, extract_pdf_text


def extract_document_text(file_bytes: bytes, filename: str) -> str:
    """Extract searchable text from PDF or PPTX bytes using the file extension."""
    suffix = Path(filename).suffix.casefold()
    if suffix == ".pdf":
        return extract_pdf_text(file_bytes)
    if suffix != ".pptx":
        raise ValueError("Unsupported proposal format. Upload a searchable PDF or PPTX file.")
    if not file_bytes:
        raise PDFExtractionError("The uploaded PowerPoint is empty.")
    try:
        presentation = Presentation(BytesIO(file_bytes))
        sections: list[str] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    value = shape.text.strip()
                    if value:
                        parts.append(value)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        if any(cells):
                            parts.append(" | ".join(cells))
            if parts:
                sections.append(f"[Slide {slide_number}]\n" + "\n".join(parts))
    except Exception as exc:
        raise PDFExtractionError(f"Could not read this PowerPoint: {exc}") from exc
    text = "\n\n".join(sections).strip()
    if not text:
        raise PDFExtractionError("No selectable text was found in this PowerPoint.")
    return text
